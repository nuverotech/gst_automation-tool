from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from pathlib import Path
from typing import List, Dict, Any
from datetime import datetime

from app.utils.logger import setup_logger

logger = setup_logger(__name__)


def generate_reconciliation_report(
    reconciled_data: List[Dict[str, Any]],
    output_path: str,
    state_wise_summary: Dict[str, Any] = None
) -> str:
    """
    Generate Excel report from reconciled GSTR2B data.
    
    Args:
        reconciled_data: List of reconciled records with comment field
        output_path: Path where to save the Excel file
        state_wise_summary: Optional state-wise summary data
        
    Returns:
        Path to generated Excel file
    """
    logger.info(f"Generating reconciliation report with {len(reconciled_data)} records")
    
    wb = Workbook()
    ws = wb.active
    ws.title = "Reconciliation Report"
    
    # Define headers
    headers = [
        "GSTIN",
        "Invoice Number",
        "Invoice Date",
        "Taxable Value",
        "POS State",
        "Comment"
    ]
    
    # Style definitions
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")
    
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Write headers
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = header
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    
    # Set column widths
    column_widths = {
        'A': 18,  # GSTIN
        'B': 20,  # Invoice Number
        'C': 15,  # Invoice Date
        'D': 15,  # Taxable Value
        'E': 12,  # POS State
        'F': 20   # Comment
    }
    
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # Write data rows
    row_idx = 2
    for record in reconciled_data:
        ws.cell(row=row_idx, column=1).value = record.get("gstin", "")
        ws.cell(row=row_idx, column=2).value = record.get("invoice_no", "")
        
        # Format date
        invoice_date = record.get("invoice_date", "")
        if isinstance(invoice_date, datetime):
            ws.cell(row=row_idx, column=3).value = invoice_date.strftime("%d-%m-%Y")
        else:
            ws.cell(row=row_idx, column=3).value = str(invoice_date)
        
        # Taxable value with 2 decimal places
        taxable_value = record.get("taxable_value", 0)
        ws.cell(row=row_idx, column=4).value = taxable_value
        ws.cell(row=row_idx, column=4).number_format = '0.00'
        
        # POS State
        pos_state = record.get("pos_state", "")
        ws.cell(row=row_idx, column=5).value = pos_state
        
        # Comment with conditional formatting
        comment = record.get("comment", "")
        comment_cell = ws.cell(row=row_idx, column=6)
        comment_cell.value = comment
        
        # Color code based on comment type
        if comment == "Matched":
            comment_cell.fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
            comment_cell.font = Font(color="006100")
        elif comment == "Not Matched":
            comment_cell.fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
            comment_cell.font = Font(color="9C6500")
        elif comment in ["Not in Books", "Not Found in 2B"]:
            comment_cell.fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
            comment_cell.font = Font(color="9C0006")
        
        # Apply borders to all cells
        for col_idx in range(1, 7):
            ws.cell(row=row_idx, column=col_idx).border = border
            ws.cell(row=row_idx, column=col_idx).alignment = Alignment(vertical="center")
        
        row_idx += 1
    
    # Add summary sheet if provided
    if state_wise_summary:
        ws_summary = wb.create_sheet("Summary")
        ws_summary.cell(row=1, column=1).value = "State-wise Reconciliation Summary"
        ws_summary.cell(row=1, column=1).font = Font(bold=True, size=14)
        
        # Headers for summary
        summary_headers = ["State", "Total Rows", "Matched", "Not Matched", "Not in Books", "Not Found in 2B"]
        for col_idx, header in enumerate(summary_headers, start=1):
            cell = ws_summary.cell(row=3, column=col_idx)
            cell.value = header
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = border
        
        # State-wise data
        row_idx = 4
        if "state_wise" in state_wise_summary:
            for state, stats in state_wise_summary["state_wise"].items():
                ws_summary.cell(row=row_idx, column=1).value = state
                ws_summary.cell(row=row_idx, column=2).value = stats.get("total_rows", 0)
                ws_summary.cell(row=row_idx, column=3).value = stats.get("matched", 0)
                ws_summary.cell(row=row_idx, column=4).value = stats.get("not_matched", 0)
                ws_summary.cell(row=row_idx, column=5).value = stats.get("not_in_books", 0)
                ws_summary.cell(row=row_idx, column=6).value = stats.get("not_in_2b", 0)
                
                for col_idx in range(1, 7):
                    ws_summary.cell(row=row_idx, column=col_idx).border = border
                
                row_idx += 1
        
        # Overall summary
        if "overall" in state_wise_summary:
            row_idx += 1
            ws_summary.cell(row=row_idx, column=1).value = "OVERALL"
            ws_summary.cell(row=row_idx, column=1).font = Font(bold=True)
            
            overall = state_wise_summary["overall"]
            ws_summary.cell(row=row_idx, column=2).value = sum(overall.values())
            ws_summary.cell(row=row_idx, column=3).value = overall.get("matched", 0)
            ws_summary.cell(row=row_idx, column=4).value = overall.get("not_matched", 0)
            ws_summary.cell(row=row_idx, column=5).value = overall.get("not_in_books", 0)
            ws_summary.cell(row=row_idx, column=6).value = overall.get("not_in_2b", 0)
            
            for col_idx in range(1, 7):
                cell = ws_summary.cell(row=row_idx, column=col_idx)
                cell.border = border
                cell.fill = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    
    # Save workbook
    output_file = Path(output_path)
    output_file.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    
    logger.info(f"Reconciliation report saved to {output_path}")
    return output_path
